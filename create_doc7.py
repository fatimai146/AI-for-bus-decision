import re
import unicodedata
import zipfile

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = "Reflection_Paper_Schiff_GRAIL.docx"

TITLE = "Reflection on Drs. Daniel and Kaylyn Schiff's Lecture on AI Governance"

SECTIONS = [
    (
        "Implications for Education",
        [
            "The result I keep thinking about is the one that failed. Kaylyn Schiff said they "
            "showed half of their respondents a video explaining what AI auditing involves, "
            "expecting that a low information space would shift once people knew more. The "
            "videos, in English and German, changed almost nothing. She read that as a counter "
            "to policymakers who claim the public is not informed enough to join "
            "decisions about AI governance. I think it cuts at educators too. We assume that "
            "teaching people about a technology changes what they want from it. Here it did not.",

            "So what is education for here? Not forming public preferences, I think, but "
            "training the people who act on them. Daniel Schiff mentioned in "
            "passing that they have been building AI auditing courses, and that is the piece my "
            "program is missing. I am taught to build and evaluate models. Policy courses "
            "teach regulation. Almost nothing teaches the middle: how you would design an audit, "
            "what evidence it should gather, who should be allowed to perform it. The six audit "
            "attributes they tested are choices somebody has to make, and they said no set "
            "standard exists yet.",

            "I am not comfortable with the generous reading of that null result. One of "
            "them described students worrying about their jobs and their majors and about AI use "
            "in educational settings, and said there is more precarity now. Stable survey "
            "preferences and anxious students are not the same public. Maybe the videos moved "
            "nobody because people had settled into positions they would not revisit, which is less flattering than principled consistency.",
        ],
    ),
    (
        "Implications for Work, Business, and the Economy",
        [
            "The headline was good news, and mostly it is. Across 3,002 adults in the "
            "United States and Germany, all three commitments they tested, public principles, an "
            "internal responsible AI team, and an independent audit, raised trust, perceived "
            "ethicality, safety, and willingness to use a company's products. Daniel Schiff put "
            "the largest effects near half a standard deviation, larger than they expected. "
            "What should worry anyone spending money is which signal earned the return. "
            "Teams and audits did not outperform a public statement of principles, and their own "
            "slide calls this the caveat that cheap signaling might still work. If I am "
            "the analyst arguing for a governance budget, that says the market pays for the press "
            "release, not the expensive work behind it. Support for regulation also barely "
            "moved, so self governing your way out of oversight does not pay either.",

            "Their interviews show where the money comes from. Among 34 auditors and "
            "governance leads across seven countries, regulatory motives and risks came out on top at "
            "104, reputational at 69, other risks at 38, and financial motives at only 22. Governance "
            "spending follows regulation. So when one of them described Anthropic rolling out text "
            "watermarks only in the European Union to satisfy the transparency code of practice, "
            "with an opt in elsewhere, it read less like an edge case than a pattern. They used "
            "the phrase strategic adjustment for that kind of move. Compliance is a "
            "map, and firms draw it as narrowly as definitions allow.",

            "Ethics did move money once. After Anthropic rejected the Pentagon's terms on "
            "February 27, 2026, Claude visits rose about 30 percent on desktop and 38 percent on "
            "mobile in their conservative estimate, with other specifications as high as 109 "
            "and 149 percent, while ChatGPT use stayed largely unchanged. But frontier firms are not ordinary "
            "companies, and they called that evidence suggestive, not definitive.",
        ],
    ),
    (
        "Implications for Me, Other People, and Society",
        [
            "Two findings sit oddly together. The public rewarded commitments as much in "
            "manufacturing and ecommerce as in finance and health care, so it did not track the "
            "risk distinction regulators built their structure around. Yet when the audits "
            "experiment varied the use case, concern ran highest for layoffs and lowest for spam "
            "filtering. People do have a line. It is not sector. It is whether the system is "
            "pointed at someone's livelihood. "
            "That makes me doubt consumer pressure as a governance mechanism. Markets respond to "
            "what people can see, and most AI harms are not visible. Their numbers say it "
            "plainly: a recent scandal cost a company roughly 25 points, while putting a "
            "government agency in charge of the audit gained roughly 9 over a self audit. Getting caught is punished far "
            "more than doing the work is rewarded, which a rational firm reads as a reputation "
            "problem, not an engineering one.",

            "And the auditor the public most wants is the one least likely to show up. External "
            "auditors beat company self audits, with government agencies ranked highest. But when "
            "an audience member asked whether AI governance would ever get serious enough for a "
            "global pact, one of the presenters said they were cautiously pessimistic about strong "
            "international governance, comparing it to climate change, where committees and voluntary "
            "agreements pile up while emissions keep rising. The research describes a public demand with no reliable supplier, and consumers "
            "cannot close that gap.",
        ],
    ),
    (
        "Personal Takeaways and Highlights",
        [
            "I came expecting a policy talk and left thinking about my own work history. The use "
            "case the public worried about most, AI used for workforce layoffs, is where I have "
            "worked. I was in human resources at Nestle and then in talent analytics at "
            "Cloudflare, so the data and models I have been near were pointed at people: who gets "
            "hired, who gets considered, how a workforce gets described to those deciding its "
            "future. I knew that was sensitive. I had not seen it at the top of a "
            "ranked list of public concern.",

            "A slide near the back of their deck, one they never reached, sharpened that. Public "
            "sector human resources managers placed more weight on transparency, privacy, and "
            "human oversight, while managers in both sectors still cared about cost, integration, "
            "bias, and consent. That second list is the conversation I remember from talent work, "
            "where constraints are concrete and tradeoffs happen fast.",

            "The highlight was Daniel Schiff describing the time he spent working on responsible "
            "AI at JPMorgan, where much of the work was case making: talking to teams, building "
            "momentum, sometimes starting from a data scientist worried about algorithmic "
            "bias. That is my likely position. Not the person writing regulation, not "
            "the executive announcing principles, but the analyst who notices something and decides "
            "whether to raise it. I used to think governance happened to a model after it "
            "was built. This moved it into the building, and left me one question. "
            "The first experiment measured willingness to use rather than actual use. Commitments "
            "changed what people said; one sharp event changed what they did. Since I "
            "will be building the thing being judged, I want to know which is the real signal.",
        ],
    ),
]

FORBIDDEN = [
    "-",
    "\u2010",
    "\u2011",
    "\u2012",
    "\u2013",
    "\u2014",
    "\u2015",
    "\u2212",
    "\u00ad",
    "\uff0d",
]


def style_run(run):
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)


def configure(paragraph, indent=True):
    pf = paragraph.paragraph_format
    pf.line_spacing = 1.15
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    if indent:
        pf.first_line_indent = Inches(0.5)


def build():
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure(title, indent=False)
    run = title.add_run(TITLE)
    run.bold = True
    style_run(run)

    for heading, paragraphs in SECTIONS:
        head = doc.add_paragraph()
        configure(head, indent=False)
        run = head.add_run(heading)
        run.bold = True
        style_run(run)

        for text in paragraphs:
            body = doc.add_paragraph()
            configure(body)
            style_run(body.add_run(text))

    doc.save(OUTPUT)


def scan_dashes(label, text):
    for char in FORBIDDEN:
        assert char not in text, "found forbidden character %r in %s" % (char, label)
    for char in set(text):
        assert unicodedata.category(char) != "Pd", (
            "found Unicode dash punctuation %r (U+%04X) in %s"
            % (char, ord(char), label)
        )


def verify():
    doc = Document(OUTPUT)
    texts = [p.text for p in doc.paragraphs]
    scan_dashes("extracted paragraph text", "\n".join(texts))

    with zipfile.ZipFile(OUTPUT) as archive:
        xml = archive.read("word/document.xml").decode("utf8")
    nodes = "".join(re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.S))
    scan_dashes("word/document.xml text nodes", nodes)

    body_words = 0
    headings = {h for h, _ in SECTIONS}
    for text in texts:
        if text and text != TITLE and text not in headings:
            body_words += len(text.split())

    print("no hyphen or dash check: PASSED")
    print("body word count:", body_words)


if __name__ == "__main__":
    build()
    verify()
