from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

GREEN = RGBColor(0x4F, 0x7F, 0x5C)
DARK = RGBColor(0x1F, 0x2A, 0x30)
BODY = RGBColor(0x33, 0x33, 0x33)
GREY = RGBColor(0x6B, 0x6B, 0x6B)
PANEL = RGBColor(0xF2, 0xF3, 0xF2)
LINE = RGBColor(0xD6, 0xD9, 0xD6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ORANGE = RGBColor(0xE0, 0x5A, 0x2B)
TEAL = RGBColor(0x2E, 0x8B, 0x57)

PHOTO = "/opt/cursor/artifacts/assets/checkfirst_server_asks.jpg"
NAME = "CheckFirst"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def text(slide, x, y, w, h, runs, size=14, color=BODY, bold=False, font="Calibri",
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Inches(0.02))
    if isinstance(runs, str):
        runs = [[(runs, {})]]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if isinstance(para, str):
            para = [(para, {})]
        for t, opts in para:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = opts.get("font", font)
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = opts.get("color", color)
            if spacing or opts.get("spacing"):
                r._r.get_or_add_rPr().set("spc", str(opts.get("spacing", spacing)))
        if "after" in (para[0][1] if para else {}):
            p.space_after = Pt(para[0][1]["after"])
    return tb


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    return s


def header(slide, label, title, page):
    text(slide, 0.67, 0.38, 10, 0.35, label, size=12, color=GREEN, bold=True, font="Arial")
    text(slide, 0.67, 0.68, 12, 0.8, title, size=32, color=DARK, bold=True, font="Arial")
    text(slide, 12.3, 6.95, 0.5, 0.3, str(page), size=10, color=GREY, align=PP_ALIGN.RIGHT)


def line(slide, x1, y1, x2, y2, color=GREY, width=1.25):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def arrow(slide, x1, y1, x2, y2):
    c = line(slide, x1, y1, x2, y2)
    ln = c.line._get_or_add_ln()
    for tag in ("a:headEnd", "a:tailEnd"):
        el = ln.makeelement(qn(tag), {"type": "triangle", "w": "med", "len": "med"})
        ln.append(el)


# Slide 1: idea
s = prs.slides.add_slide(BLANK)
header(s, "IDEA 2 OF 3  |  STANDARD SPLIT QUESTION",
       f"{NAME}: the restaurant asks, so no friend has to", 7)
text(s, 0.67, 1.6, 12, 0.5,
     [[("Need  ", {"bold": True, "color": GREEN}),
       ("Asking for separate checks sounds cheap when a friend brings it up, but it is just logistics when the restaurant asks every table.", {})]],
     size=15)

rect(s, 0.67, 2.3, 6.75, 4.45, PANEL)
s.shapes.add_picture(PHOTO, Inches(0.85), Inches(2.5), Inches(3.75), Inches(2.81))
text(s, 0.85, 5.4, 3.75, 1.2,
     [[("Asked at the start of every meal, ", {"bold": True, "color": DARK}),
       ("as routinely as \"still or sparkling water?\"", {"italic": True})]],
     size=12)

# handheld mockup
px, py, pw, ph = 4.85, 2.5, 2.4, 4.05
rect(s, px, py, pw, ph, DARK, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
rect(s, px + 0.12, py + 0.2, pw - 0.24, ph - 0.4, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, px + 0.25, py + 0.32, pw - 0.5, 0.3, "Table 12  ·  New order", size=10, color=GREY)
text(s, px + 0.25, py + 0.6, pw - 0.5, 0.6, "How is this table paying?", size=13, color=DARK, bold=True, font="Arial")
opts = ["One check", "Split evenly", "Separate checks", "Decide later"]
for i, o in enumerate(opts):
    b = rect(s, px + 0.25, py + 1.3 + i * 0.5, pw - 0.5, 0.38, WHITE if i else GREEN,
             MSO_SHAPE.ROUNDED_RECTANGLE, line=GREEN, radius=0.3)
    tf = b.text_frame
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = o
    r.font.size = Pt(11)
    r.font.name = "Calibri"
    r.font.bold = True
    r.font.color.rgb = WHITE if i == 0 else GREEN
rect(s, px + 0.25, py + 3.35, pw - 0.5, 0.34, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
text(s, px + 0.25, py + 3.35, pw - 0.5, 0.34, "Pick one to send the first order",
     size=9, color=GREY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, italic=True)

steps = [
    ("1  Asked at every table",
     "When the server takes the drink order, they ask \"One check, split evenly, or separate checks?\" Every table hears it, every time."),
    ("2  Built into the order",
     "The server's handheld will not send the first order until a check style is picked, so no table gets skipped."),
    ("3  The bill arrives ready",
     "Separate checks are already set up when the meal ends. Nobody at the table had to be the one to bring it up."),
]
y = 2.4
for h, b in steps:
    text(s, 7.75, y, 5.0, 0.35, h, size=15, color=DARK, bold=True)
    text(s, 7.75, y + 0.36, 5.0, 0.9, b, size=13.5)
    y += 1.4

# Slide 2: positioning statement
s = prs.slides.add_slide(BLANK)
header(s, f"IDEA 2 OF 3  |  {NAME.upper()}", "Positioning statement", 8)
rows = [
    ("FOR", "friends who eat out together on different budgets"),
    ("WHO", "accept an even split because being the one to ask for separate checks feels cheap"),
    (f"{NAME.upper()} IS", "a restaurant service routine built into the server's ordering step"),
    ("THAT", "asks every table how they are paying before anyone orders"),
    ("UNLIKE", "waiting for a friend to ask the server when the bill comes, or splitting it later on Splitwise or Venmo"),
    ("OUR PRODUCT", "makes the split a question the restaurant asks everyone, so no friend has to raise it"),
]
y0, rh = 1.75, 0.76
for i, (k, v) in enumerate(rows):
    y = y0 + i * rh
    text(s, 0.67, y, 1.65, rh, k, size=13, color=GREEN, bold=True, font="Arial",
         anchor=MSO_ANCHOR.MIDDLE, spacing=100)
    text(s, 2.45, y, 6.2, rh, v, size=14, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(rows) - 1:
        line(s, 0.67, y + rh, 8.7, y + rh, LINE, 0.75)
box = rect(s, 9.25, 1.75, 3.5, 4.55, GREEN, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, 9.55, 2.05, 3.0, 0.4, "THE ONE BENEFIT", size=12, color=WHITE, bold=True, font="Arial", spacing=400)
text(s, 9.55, 2.55, 3.0, 3.5, "Separate checks become a normal choice, not a personal request.",
     size=24, color=WHITE, bold=True, font="Arial")

# Slide 3: perceptual map
s = prs.slides.add_slide(BLANK)
header(s, f"IDEA 2 OF 3  |  {NAME.upper()}", "Perceptual map", 9)
cx, cy = 4.5, 4.2
left, right, top, bottom = 0.8, 8.25, 2.15, 6.3
rect(s, cx, top, right - cx, cy - top, PANEL)
arrow(s, left, cy, right, cy)
arrow(s, cx, top, cx, bottom)
text(s, cx - 2, top - 0.42, 4, 0.35, "Every table, every time", size=13, color=DARK, bold=True,
     font="Arial", align=PP_ALIGN.CENTER)
text(s, cx - 2, bottom + 0.05, 4, 0.35, "Only if someone brings it up", size=13, color=DARK, bold=True,
     font="Arial", align=PP_ALIGN.CENTER)
text(s, left - 0.05, cy + 0.08, 2.0, 0.6, ["A friend has", "to raise it"], size=13, color=DARK,
     bold=True, font="Arial")
text(s, right - 2.0, cy + 0.08, 2.0, 0.6, ["The restaurant", "raises it"], size=13, color=DARK,
     bold=True, font="Arial", align=PP_ALIGN.RIGHT)


def pill(slide, x, y, w, label, h=0.48):
    p = rect(slide, x, y, w, h, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=RGBColor(0xB5, 0xB5, 0xB5), radius=0.5)
    tf = p.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run()
    r.text = label
    r.font.size = Pt(11)
    r.font.name = "Calibri"
    r.font.color.rgb = BODY


pill(s, 1.3, 2.65, 2.4, "One card, split evenly")
pill(s, 1.0, 4.85, 2.6, "Asking the server for separate checks", h=0.6)
pill(s, 1.75, 5.6, 2.1, "Splitwise / Venmo later")
pill(s, 5.15, 4.75, 2.6, "Some servers ask \"together or separate?\"", h=0.6)
pill(s, 5.5, 5.5, 2.3, "QR pay at the table")

hero = rect(s, 5.55, 2.7, 2.2, 0.85, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE, line=GREEN, radius=0.2)
hero.line.width = Pt(2)
tf = hero.text_frame
tf.paragraphs[0].alignment = PP_ALIGN.CENTER
r = tf.paragraphs[0].add_run()
r.text = NAME
r.font.size = Pt(20)
r.font.bold = True
r.font.name = "Arial"
r.font.color.rgb = GREEN

text(s, 8.9, 1.9, 4.0, 0.4, "Why these axes", size=16, color=DARK, bold=True, font="Arial")
text(s, 8.9, 2.4, 4.0, 2.2,
     [[("•  X: who carries the social cost of raising the split.", {"after": 8})],
      [("•  Y: consistency. A question only feels like normal logistics if every table gets it.", {})]],
     size=13)
rect(s, 8.75, 4.55, 4.25, 1.95, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 8.95, 4.72, 3.85, 1.65,
     [[("Takeaway: ", {"bold": True, "color": GREEN}),
       ("Some servers already ask, but only sometimes, so asking still looks unusual. "
        f"{NAME} asks every table, every time, which is what turns the split into logistics.", {})]],
     size=13)

# Slide 4: reverse interviewing (replacement for slide 13)
s = prs.slides.add_slide(BLANK)
text(s, 0.67, 0.38, 10, 0.35, "REVERSE INTERVIEWING", size=12, color=GREY, bold=True, font="Arial")
text(s, 0.67, 0.68, 12, 0.8, "What an LLM asked us and how we answered", size=32, color=DARK,
     bold=True, font="Arial")
text(s, 12.3, 6.95, 0.5, 0.3, "13", size=10, color=GREY, align=PP_ALIGN.RIGHT)
cols = [
    ("SplitSeat", ORANGE, [
        ("How do shared appetizers get charged?",
         "Servers tag shared items to \"table,\" and those split evenly across everyone."),
        ("Why would restaurants adopt it?",
         "Faster table turns: no server time spent running five cards."),
        ("What if a server forgets to tag a seat?",
         "The POS asks for a seat number before an item can be sent to the kitchen."),
    ]),
    (NAME, GREEN, [
        ("Won't the person who picks \"separate\" still feel singled out?",
         "The server asks the whole table, with separate checks as an ordinary option, like still or sparkling."),
        ("What if the group is not sure yet?",
         "\"Decide later\" is an option, and the answer can change with one tap before the bill prints."),
        ("Why would restaurants adopt it?",
         "It costs nothing to start, and servers stop rerunning bills when a table asks to split at the end."),
    ]),
    ("SplitSense", TEAL, [
        ("In a group of two, is it really anonymous?",
         "No, so it activates only at tables of three or more."),
        ("Why would a diner bother setting a limit?",
         "It is one tap at check-in with a smart default already set."),
        (f"How is this different from {NAME}?",
         f"{NAME} makes the question normal for everyone. SplitSense protects people who still would not say it."),
    ]),
]
for i, (name, col, qa) in enumerate(cols):
    x = 0.67 + i * 4.15
    rect(s, x, 1.75, 3.85, 5.0, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    paras = [[(name, {"size": 17, "bold": True, "color": col, "font": "Arial", "after": 10})]]
    for q, a in qa:
        paras.append([("Q: " + q, {"bold": True, "color": DARK, "after": 3})])
        paras.append([("A: " + a, {"after": 12})])
    text(s, x + 0.2, 1.95, 3.45, 4.7, paras, size=12)

prs.save("/workspace/CheckFirst_Standard_Split_Question_Slides.pptx")
print("saved")
